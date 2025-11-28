"""PPT 生成服务"""
import logging
import os
from pathlib import Path
from typing import Dict, Any, Optional
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from PIL import Image

logger = logging.getLogger(__name__)


class PPTService:
    """PPT 生成服务类"""

    def __init__(self):
        logger.debug("初始化 PPTService...")
        self.history_root_dir = Path(__file__).parent.parent.parent / "history"
        self.template_config = self._load_template_config()
        logger.info("PPTService 初始化完成")

    def _load_template_config(self) -> dict:
        """加载 PPT 模板配置"""
        config_path = Path(__file__).parent.parent.parent / 'ppt_templates.yaml'
        if config_path.exists():
            import yaml
            with open(config_path, 'r', encoding='utf-8') as f:
                return yaml.safe_load(f) or {}

        # 默认配置
        return {
            'slide_width': 10,      # 英寸
            'slide_height': 5.625,  # 英寸（16:9）
            'title_font_size': 44,
            'content_font_size': 24,
            'title_font_name': '微软雅黑',
            'content_font_name': '微软雅黑',
        }

    def generate_ppt(
        self,
        task_id: str,
        outline: Dict[str, Any],
        template: str = 'default'
    ) -> str:
        """
        生成 PPT 文件

        Args:
            task_id: 任务 ID
            outline: 大纲数据 (包含 pages 列表)
            template: 模板名称

        Returns:
            PPT 文件路径
        """
        try:
            logger.info(f"开始生成 PPT: task_id={task_id}, template={template}")

            # 创建 Presentation 对象
            prs = Presentation()
            prs.slide_width = Inches(self.template_config['slide_width'])
            prs.slide_height = Inches(self.template_config['slide_height'])

            # 任务目录
            task_dir = self.history_root_dir / task_id
            if not task_dir.exists():
                raise FileNotFoundError(f"任务目录不存在: {task_id}")

            pages = outline.get('pages', [])
            logger.debug(f"大纲包含 {len(pages)} 页")

            # 遍历页面，生成幻灯片
            for page in pages:
                index = page['index']
                page_type = page['type']
                content = page['content']

                # 图片路径
                image_path = task_dir / f"{index}.png"
                if not image_path.exists():
                    logger.warning(f"图片不存在: {image_path}，跳过该页")
                    continue

                # 添加空白幻灯片
                blank_slide_layout = prs.slide_layouts[6]  # 空白布局
                slide = prs.slides.add_slide(blank_slide_layout)

                # 添加图片（全屏或居中）
                self._add_image_to_slide(slide, image_path, prs)

                # 可选：添加文本框（覆盖在图片上）
                if page_type == 'cover':
                    # 封面：只添加主标题
                    self._add_title_textbox(slide, content, prs)

            # 保存 PPT
            output_filename = f"{task_id}.pptx"
            output_path = task_dir / output_filename
            prs.save(str(output_path))

            logger.info(f"✅ PPT 生成成功: {output_path}")
            return str(output_path)

        except Exception as e:
            logger.error(f"❌ PPT 生成失败: {str(e)}")
            raise

    def _add_image_to_slide(self, slide, image_path: Path, prs: Presentation):
        """
        将图片添加到幻灯片（自适应缩放，保持比例）

        Args:
            slide: 幻灯片对象
            image_path: 图片路径
            prs: Presentation 对象
        """
        # 获取图片尺寸
        img = Image.open(image_path)
        img_width, img_height = img.size
        img_ratio = img_width / img_height

        # 幻灯片尺寸
        slide_width = prs.slide_width
        slide_height = prs.slide_height
        slide_ratio = slide_width / slide_height

        # 计算缩放后的尺寸（保持比例，填充幻灯片）
        if img_ratio > slide_ratio:
            # 图片更宽，以高度为准
            pic_height = slide_height
            pic_width = int(pic_height * img_ratio)
        else:
            # 图片更高，以宽度为准
            pic_width = slide_width
            pic_height = int(pic_width / img_ratio)

        # 居中位置
        left = (slide_width - pic_width) // 2
        top = (slide_height - pic_height) // 2

        # 添加图片
        slide.shapes.add_picture(
            str(image_path),
            left, top,
            width=pic_width,
            height=pic_height
        )

    def _add_title_textbox(self, slide, text: str, prs: Presentation):
        """
        在幻灯片顶部添加标题文本框

        Args:
            slide: 幻灯片对象
            text: 标题文本
            prs: Presentation 对象
        """
        # 提取标题（取第一行或前50字）
        title_text = text.split('\n')[0][:50]

        # 添加文本框（顶部居中）
        left = Inches(1)
        top = Inches(0.5)
        width = prs.slide_width - Inches(2)
        height = Inches(1)

        textbox = slide.shapes.add_textbox(left, top, width, height)
        text_frame = textbox.text_frame
        text_frame.text = title_text

        # 样式设置
        paragraph = text_frame.paragraphs[0]
        paragraph.alignment = PP_ALIGN.CENTER
        paragraph.font.size = Pt(self.template_config['title_font_size'])
        paragraph.font.name = self.template_config['title_font_name']
        paragraph.font.bold = True

    def get_ppt_path(self, task_id: str) -> Optional[str]:
        """
        获取已生成的 PPT 路径

        Args:
            task_id: 任务 ID

        Returns:
            PPT 文件路径，如果不存在则返回 None
        """
        ppt_path = self.history_root_dir / task_id / f"{task_id}.pptx"
        if ppt_path.exists():
            return str(ppt_path)
        return None


# 全局服务实例
_ppt_service_instance = None


def get_ppt_service() -> PPTService:
    """获取 PPT 生成服务实例"""
    global _ppt_service_instance
    if _ppt_service_instance is None:
        _ppt_service_instance = PPTService()
    return _ppt_service_instance
